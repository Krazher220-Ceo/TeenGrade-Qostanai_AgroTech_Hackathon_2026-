"""Сборка PPTX-презентации питча (3 мин) с встроенными видео.

Запуск: .venv/bin/python scripts/build_pitch_pptx.py
Результат: docs/AgroVision_pitch.pptx
"""
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "AgroVision_pitch.pptx"
ANIM = ROOT / "design" / "dock_station_animation.mp4"
SIM = ROOT / "case1" / "output" / "dock_simulation" / "dock_mission_simulation.mp4"
ASSETS = ROOT / "docs" / "assets"
MOCK = ROOT / "design" / "mockups"

DARK = RGBColor(0x0F, 0x3D, 0x2E)
GREEN = RGBColor(0x1F, 0x9D, 0x55)
TEXT = RGBColor(0x1E, 0x29, 0x24)
MUTED = RGBColor(0x5B, 0x6B, 0x63)
AMBER = RGBColor(0xE8, 0x9B, 0x0C)
BG = RGBColor(0xF6, 0xF8, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Emu(12192000), Emu(6858000)  # 16:9
IN = 914400

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
TMP = Path(tempfile.mkdtemp())


def poster(video: Path, t: float) -> str:
    out = TMP / f"{video.stem}_{t}.png"
    subprocess.run(["ffmpeg", "-loglevel", "error", "-y", "-ss", str(t), "-i", str(video),
                    "-frames:v", "1", str(out)], check=True)
    return str(out)


def bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def text(slide, x, y, w, h, s, size=18, color=TEXT, bold=False, align=None):
    tb = slide.shapes.add_textbox(Emu(int(x * IN)), Emu(int(y * IN)), Emu(int(w * IN)), Emu(int(h * IN)))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Inter"
        p.space_after = Pt(size * 0.45)
        if align is not None:
            p.alignment = align
    return tb


def header(slide, kicker, title):
    bar = slide.shapes.add_shape(1, 0, 0, Emu(int(0.18 * IN)), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()
    text(slide, 0.6, 0.35, 12, 0.4, kicker.upper(), 13, GREEN, True)
    text(slide, 0.6, 0.7, 12, 0.9, title, 32, DARK, True)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def bullets(slide, x, y, w, items, size=20):
    text(slide, x, y, w, 5, ["•  " + i for i in items], size, TEXT)


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Emu(int(w * IN))
    if h:
        kw["height"] = Emu(int(h * IN))
    return slide.shapes.add_picture(str(path), Emu(int(x * IN)), Emu(int(y * IN)), **kw)


def video(slide, path, x, y, w, h, t_poster):
    return slide.shapes.add_movie(str(path), Emu(int(x * IN)), Emu(int(y * IN)),
                                  Emu(int(w * IN)), Emu(int(h * IN)),
                                  poster_frame_image=poster(path, t_poster), mime_type="video/mp4")


# 1. Обложка
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
text(s, 0.9, 2.1, 11, 1.2, "AgroVision AI", 60, WHITE, True)
text(s, 0.9, 3.25, 11, 1.2, "Карта сорняков за один выезд агронома: где, какого вида, в какой фазе и что делать",
     24, RGBColor(0xC9, 0xE8, 0xD4))
text(s, 0.9, 5.6, 11, 0.5, "TEAM1 · Кейс №1 «Олжа Агро» · Qostanai AgroTech Hackathon 2026", 16,
     RGBColor(0x9F, 0xC9, 0xB0))
notes(s, "Здравствуйте, мы команда TEAM1, проект AgroVision AI. Мы делаем так, чтобы по снимкам с дрона "
         "агроном за один выезд получал карту сорняков: где они, какого вида, в какой фазе и что с ними делать.")

# 2. Две модели
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Как это работает", "Две модели: найти → распознать")
for i, (t, d) in enumerate([
    ("1. Детектор", "Обучен на культурах и сорняках. На снимке с дрона обводит каждый сорняк рамкой. "
                    "Каждая рамка — отдельный фрагмент, до 80 на кадр."),
    ("2. Классификатор", "По фрагменту определяет вид (26 видов из датасета кейса) и фазу развития."),
    ("3. Правила ментора", "Превращают вид, фазу и плотность в решение: обрабатывать или нет, какой дозой."),
]):
    x = 0.6 + i * 4.1
    box = s.shapes.add_shape(1, Emu(int(x * IN)), Emu(int(2.0 * IN)), Emu(int(3.8 * IN)), Emu(int(3.6 * IN)))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = RGBColor(0xDD, 0xE5, 0xDF)
    text(s, x + 0.25, 2.25, 3.3, 0.6, t, 24, GREEN, True)
    text(s, x + 0.25, 3.35, 3.3, 2.2, d, 18, TEXT)
notes(s, "Внутри две модели. Первая — детектор. Она обучена на культурах и сорняках и на снимке с дрона обводит "
         "каждый сорняк рамкой. Каждая рамка вырезается как отдельный фрагмент: из одного кадра их бывает до "
         "восьмидесяти. Вторая модель — классификатор. По фрагменту она определяет вид, один из 26 видов из "
         "датасета кейса, и фазу развития. Дальше правила ментора превращают это в решение.")

# 3. Результаты
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Результаты на данных кейса", "Что показывают модели")
for i, (big, small) in enumerate([
    ("0,76", "mAP@50 детектора\n313 реальных тестовых снимков"),
    ("1 / 12", "ложное срабатывание\nна кадрах чистой почвы"),
    ("98,4% · 98,7%", "классификатор: вид · фаза\n1 591 эталонное фото"),
    ("0,53 с", "на 4K-кадр\nна GPU сервера"),
]):
    x = 0.6 + (i % 2) * 6.1
    y = 1.9 + (i // 2) * 2.3
    text(s, x, y, 5.8, 1.0, big, 44, DARK, True)
    text(s, x, y + 1.0, 5.8, 1.0, small.split("\n"), 16, MUTED)
notes(s, "Детектор на реальных тестовых снимках даёт mAP@50 0,76 и почти не срабатывает на чистой почве: одна "
         "ошибка на двенадцати кадрах. Классификатор на эталонных фото: 98,4% по виду и 98,7% по фазе. Один "
         "4K-кадр сервер обрабатывает примерно за полсекунды.")

# 4. Мобильная дрон-станция — анимация
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Выезд в поле", "Мобильная дрон-станция на машине")
if ANIM.exists():
    video(s, ANIM, 0.6, 1.65, 8.8, 4.95, 8)
bullets(s, 9.7, 1.8, 3.4, [
    "Крыша открывается, 3–4 дрона взлетают по очереди на разные эшелоны",
    "Приложение делит поле на сектора",
    "Посадка → зарядка → выгрузка",
    "Детектор в машине, через Starlink — только фрагменты + GPS",
], 15)
notes(s, "Агроном выезжает на машине с док-станцией на крыше. Наше приложение делит поле между тремя-четырьмя "
         "дронами и рассчитывает высоту, угол камеры и маршруты. Дроны взлетают по очереди, облетают свои "
         "сектора, садятся обратно в станцию и заряжаются. Станция сама выгружает снимки, детектор работает "
         "прямо в машине, и через Starlink на сервер уходят только вырезанные фрагменты с координатами — "
         "сырые фото весят десятки гигабайт и через спутник не пройдут.")

# 5. Human-in-the-Loop
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Human-in-the-Loop", "Где модель не уверена — решает агроном")
bullets(s, 0.6, 1.9, 6.4, [
    "Всё ниже 75% уверенности — в приложение агронома",
    "Он видит фрагмент и варианты модели: выбирает ближайший или указывает вид сам",
    "Приложение работает без связи",
    "Каждое решение — новая разметка для дообучения",
], 20)
if (MOCK / "agrovision-android-concept.png").exists():
    picture(s, MOCK / "agrovision-android-concept.png", 7.4, 1.6, h=5.0)
notes(s, "На снимках с дрона модель не всегда уверена, и мы это не прячем. Всё, что ниже 75% уверенности, "
         "попадает в приложение агронома. Он видит фрагмент и варианты модели, выбирает ближайший или указывает "
         "вид сам. Приложение работает без связи. Каждое такое решение сохраняется и идёт на дообучение, "
         "поэтому с каждым выездом модель ошибается реже.")

# 6. Правила ментора
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Агрономия", "Решение по правилам ментора")
bullets(s, 0.6, 1.9, 12, [
    "Порог вредоносности по плотности на м² — отдельно для малолетних и многолетних",
    "Поправка дозы по фазе развития сорняка",
    "Переросший сорняк → предупреждение: обработка будет малоэффективной",
], 24)
notes(s, "Решение принимается по правилам, которые нам дал ментор. Первое — порог вредоносности по плотности на "
         "квадратный метр, отдельно для малолетних и многолетних сорняков. Второе — поправка дозы по фазе. Если "
         "сорняк уже переросший, система предупреждает агронома, что обработка будет малоэффективной.")

# 7. От очагов к заданию
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Результат", "От очагов к заданию для опрыскивателя")
bullets(s, 0.6, 1.9, 5.6, [
    "Вокруг очага — зона с GPS: рядом почти всегда есть ещё растения",
    "У многолетников зона шире, они уходят в отдельную базу для сравнения по сезонам",
    "Выгрузка карты для опрыскивателя (ISO-XML TaskData)",
    "Общий дашборд агрономов",
], 19)
if (ASSETS / "field_demo_overview.png").exists():
    picture(s, ASSETS / "field_demo_overview.png", 6.6, 1.7, w=6.1)
notes(s, "На выходе не просто точки. Вокруг каждого очага строится зона с координатами, потому что рядом почти "
         "всегда есть ещё растения, а у многолетников зона шире. Зоны выгружаются картой для опрыскивателя. "
         "Многолетники дополнительно сохраняются в отдельную базу, чтобы сравнивать поле от сезона к сезону. "
         "Всё это видно в общем дашборде агрономов.")

# 8. Демо — симуляция облёта
s = prs.slides.add_slide(BLANK)
bg(s)
header(s, "Демо", "Симуляция облёта: 4 дрона из одной машины")
if SIM.exists():
    video(s, SIM, 1.2, 1.6, 10.9, 10.9 * 9 / 16 * 0.93, 12)
notes(s, "Короткое видео: машина встаёт на край поля в точке, которую подсказал планировщик, четыре дрона "
         "облетают свои сектора, возвращаются на станцию, данные уходят на сервер и в очередь агронома.")

# 9. Итог
s = prs.slides.add_slide(BLANK)
bg(s, DARK)
text(s, 0.9, 1.8, 11.5, 1, "Дрон снимает → две модели находят сорняки и определяют вид и фазу → "
     "агроном подтверждает сомнительное → готовая карта обработки", 30, WHITE, True)
text(s, 0.9, 4.6, 11, 0.8, "Спасибо! Готовы ответить на вопросы", 26, RGBColor(0xC9, 0xE8, 0xD4))
notes(s, "Итог: дрон снимает, две модели находят сорняки и определяют вид и фазу, агроном подтверждает "
         "сомнительное, и на выходе готовая карта обработки. Спасибо, готовы ответить на вопросы.")

prs.save(OUT)
print(OUT)
